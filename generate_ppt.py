import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # completely blank layout

    # Color Palette
    BG_DARK = RGBColor(11, 19, 43)        # #0B132B Deep Navy
    BG_CARD = RGBColor(28, 37, 65)        # #1C2541 Slate Card
    BG_CARD_LIGHT = RGBColor(38, 50, 85)  # Slightly lighter card
    CYAN_ACCENT = RGBColor(0, 212, 255)   # #00D4FF Bright Cyan
    PURPLE_ACCENT = RGBColor(168, 85, 247)# #A855F7 Vibrant Purple
    GREEN_ACCENT = RGBColor(34, 197, 94)  # #22C55E Emerald Green
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(156, 163, 175)  # #9CA3AF Gray
    BORDER_COLOR = RGBColor(51, 65, 85)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="NETSCAN PRO"):
        # Category / Kicker
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = CYAN_ACCENT

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

    def add_card(slide, left, top, width, height, bg_color=BG_CARD, border_color=BORDER_COLOR):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
        return card

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Decorative hero card
    add_card(slide1, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5), BG_CARD, CYAN_ACCENT)

    # Title box
    tb = slide1.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.333), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "NETWORK SECURITY & INFRASTRUCTURE"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = CYAN_ACCENT
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "NetScan Pro"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "Real-Time Network Device Scanner & Access Management System"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(226, 232, 240)
    p2.space_after = Pt(28)

    p3 = tf.add_paragraph()
    p3.text = "Comprehensive Guide: Need, Architecture, Significance & Operation"
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 2: Problem Statement & The Need
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "The Problem: The Invisible Network Challenge", "THE NEED & MOTIVATION")

    cards_data_s2 = [
        ("Network Blind Spots", "In modern homes and enterprises, dozens of devices connect to Wi-Fi. Users lack visibility into what is actively connected.", "📡", CYAN_ACCENT),
        ("Rogue & Unsecured Devices", "Unapproved IoT hardware, piggybacking neighbors, or rogue sniffers can join the network without raising any immediate alarm.", "🛡️", PURPLE_ACCENT),
        ("Troubleshooting Friction", "Diagnosing printer disconnects or server latency requires complex CLI commands (ping, arp, nmap) or router admin logins.", "⚡", GREEN_ACCENT)
    ]

    for i, (title, desc, icon, accent) in enumerate(cards_data_s2):
        x = Inches(0.8 + i * 3.95)
        add_card(slide2, x, Inches(1.8), Inches(3.75), Inches(4.8), BG_CARD, accent)
        
        tb = slide2.shapes.add_textbox(x + Inches(0.3), Inches(2.1), Inches(3.15), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_icon = tf.paragraphs[0]
        p_icon.text = icon
        p_icon.font.size = Pt(36)
        p_icon.space_after = Pt(14)
        
        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(18)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.space_after = Pt(12)
        
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 3: How It Works (5-Step Engine)
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "How NetScan Pro Operates: 5 Core Steps", "SYSTEM WORKFLOW")

    steps = [
        ("1. SSID & Network Profile Isolation", "Automatically identifies Wi-Fi SSID and adapter CIDR subnet, provisioning an isolated device namespace per network environment."),
        ("2. Batched Ping Sweeps & Concurrency", "Executes concurrency-controlled ping sweeps across all 254 IP addresses every 6 seconds to discover alive hosts and measure latency without freezing."),
        ("3. ARP Inspection & Hardware OUI Resolution", "Extracts hardware MAC addresses from the ARP cache and matches OUI prefixes to identify manufacturers (Apple, Samsung, Intel, HP)."),
        ("4. Targeted Unicast IP Probing (⚡ Probe IP)", "Allows administrators to enter custom IPs (e.g., 172.20.181.188) to probe and track cross-subnet nodes or silent clients on demand."),
        ("5. SQLite Persistence & Live WebSocket Sync", "Logs every connection/disconnection event permanently in SQLite with timestamps and broadcasts live state updates via Socket.IO.")
    ]

    for i, (stitle, sdesc) in enumerate(steps):
        y = Inches(1.75 + i * 1.05)
        add_card(slide3, Inches(0.8), y, Inches(11.733), Inches(0.9), BG_CARD)
        
        tb = slide3.shapes.add_textbox(Inches(1.1), y + Inches(0.12), Inches(11.1), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = stitle + " — "
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = CYAN_ACCENT
        
        run = p.add_run()
        run.text = sdesc
        run.font.size = Pt(13)
        run.font.bold = False
        run.font.color.rgb = TEXT_WHITE

    # ==========================================
    # SLIDE 4: Key Features & Capabilities
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Key Features & Capabilities", "FEATURE MATRIX")

    features = [
        ("🌐 Multi-Network Profile Isolation", "Automatically identifies SSIDs/subnets, isolating devices per Wi-Fi network with an active/historical profile switcher.", "🌐"),
        ("⚡ On-Demand Target Probing (Probe IP)", "Enter any specific IP to send direct ICMP/ARP probes and force-track custom or cross-subnet nodes.", "🎯"),
        ("📜 Persistent DB History & CSV Export", "Full historical audit trail in SQLite recording every connect/disconnect event with 1-click CSV download.", "💾"),
        ("🏷️ Hardware Vendor Identification", "Automatic MAC OUI dictionary matching identifies Apple, Samsung, OnePlus, Dell, HP, Xiaomi, and Google devices.", "🏷️"),
        ("🔔 Real-Time Live Toast Alerts", "Instant green/red popups when devices connect or drop off without requiring page reloads.", "🔔"),
        ("🛡️ Role-Based Access Control (RBAC)", "Secure dual-role system (Admin vs User) with bcrypt password hashing and user governance.", "🛡️")
    ]

    for i, (ftitle, fdesc, ficon) in enumerate(features):
        row = i // 3
        col = i % 3
        x = Inches(0.8 + col * 3.95)
        y = Inches(1.8 + row * 2.55)
        
        add_card(slide4, x, y, Inches(3.75), Inches(2.35), BG_CARD)
        tb = slide4.shapes.add_textbox(x + Inches(0.25), y + Inches(0.2), Inches(3.25), Inches(1.95))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = f"{ficon}  {ftitle}"
        p0.font.size = Pt(15)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_WHITE
        p0.space_after = Pt(8)
        
        p1 = tf.add_paragraph()
        p1.text = fdesc
        p1.font.size = Pt(12)
        p1.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 5: Role-Based Access Control (RBAC)
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Role-Based Access Control (RBAC)", "SECURITY & GOVERNANCE")

    # Left Card: Standard User
    add_card(slide5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), BG_CARD, CYAN_ACCENT)
    tb_u = slide5.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.2))
    tf_u = tb_u.text_frame
    tf_u.word_wrap = True

    pu0 = tf_u.paragraphs[0]
    pu0.text = "👤  Standard User Role"
    pu0.font.size = Pt(20)
    pu0.font.bold = True
    pu0.font.color.rgb = CYAN_ACCENT
    pu0.space_after = Pt(14)

    user_points = [
        "Full access to the live Network Dashboard",
        "View active & offline devices with real-time latency",
        "Perform on-demand network scans & search filters",
        "Receive live connection & disconnection toast alerts",
        "Restricted from viewing or altering system user accounts"
    ]
    for pt in user_points:
        p = tf_u.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(8)

    # Right Card: Admin User
    add_card(slide5, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8), BG_CARD, PURPLE_ACCENT)
    tb_a = slide5.shapes.add_textbox(Inches(7.2), Inches(2.1), Inches(5.0), Inches(4.2))
    tf_a = tb_a.text_frame
    tf_a.word_wrap = True

    pa0 = tf_a.paragraphs[0]
    pa0.text = "🛡️  Administrator Role"
    pa0.font.size = Pt(20)
    pa0.font.bold = True
    pa0.font.color.rgb = PURPLE_ACCENT
    pa0.space_after = Pt(14)

    admin_points = [
        "All privileges of a Standard User (Dual Access)",
        "Dedicated Admin Panel for user account management",
        "Promote or demote user roles dynamically",
        "Delete obsolete accounts and audit logins",
        "View enterprise analytics (Active 24h, total accounts, join dates)"
    ]
    for pt in admin_points:
        p = tf_a.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(8)

    # ==========================================
    # SLIDE 6: Technical Architecture & Stack
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Technical Architecture & Tech Stack", "FULL-STACK ARCHITECTURE")

    tech_data = [
        ("Frontend Layer", "React 19 + Vite", "High-performance SPA with modular components, responsive CSS Grid layout, real-time Socket.io-client bindings, and Auth context.", CYAN_ACCENT),
        ("Backend & Scanner Engine", "Express 5 + Socket.io", "Node.js REST API with batched child-process ping sweeps, reverse DNS resolution, and bidirectional WebSocket event streams.", PURPLE_ACCENT),
        ("Database & Security", "SQLite (better-sqlite3) + JWT", "Lightweight, zero-config relational database with bcrypt-hashed passwords, JWT tokens, and live role validation middleware.", GREEN_ACCENT)
    ]

    for i, (layer, stack, details, accent) in enumerate(tech_data):
        x = Inches(0.8 + i * 3.95)
        add_card(slide6, x, Inches(1.8), Inches(3.75), Inches(4.8), BG_CARD, accent)
        
        tb = slide6.shapes.add_textbox(x + Inches(0.3), Inches(2.1), Inches(3.15), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_l = tf.paragraphs[0]
        p_l.text = layer.upper()
        p_l.font.size = Pt(12)
        p_l.font.bold = True
        p_l.font.color.rgb = accent
        p_l.space_after = Pt(6)
        
        p_s = tf.add_paragraph()
        p_s.text = stack
        p_s.font.size = Pt(18)
        p_s.font.bold = True
        p_s.font.color.rgb = TEXT_WHITE
        p_s.space_after = Pt(14)
        
        p_d = tf.add_paragraph()
        p_d.text = details
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 7: Real-World Use Cases & Applications
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Real-World Use Cases & Practical Applications", "PRACTICAL VALUE")

    use_cases = [
        ("Home Security & IoT Management", "Track smart speakers, IP cameras, Smart TVs, and phones. Detect unknown or neighbor devices immediately if your Wi-Fi password is compromised.", "🏠"),
        ("Office & Small Business IT", "Maintain a live inventory of employee workstations, servers, and shared network printers without purchasing expensive enterprise network tools.", "🏢"),
        ("Network Performance & Health", "Identify dead zones, weak Wi-Fi signals, high-latency devices, and dropped packets in real time directly from any web browser.", "📊"),
        ("Lab & Development Environments", "Easily find dynamic DHCP IP addresses assigned to Raspberry Pis, microcontrollers, test servers, and VMs without looking up router tables.", "🔬")
    ]

    for i, (utitle, udesc, uicon) in enumerate(use_cases):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.8 + row * 2.55)
        
        add_card(slide7, x, y, Inches(5.75), Inches(2.35), BG_CARD)
        tb = slide7.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), Inches(5.15), Inches(1.95))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = f"{uicon}  {utitle}"
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_WHITE
        p0.space_after = Pt(8)
        
        p1 = tf.add_paragraph()
        p1.text = udesc
        p1.font.size = Pt(13)
        p1.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 8: Summary & Conclusion
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Conclusion & Future Roadmap", "SUMMARY")

    add_card(slide8, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8), BG_CARD, CYAN_ACCENT)
    tb_c = slide8.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.9), Inches(4.0))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    pc0 = tf_c.paragraphs[0]
    pc0.text = "Why NetScan Pro Stands Out"
    pc0.font.size = Pt(22)
    pc0.font.bold = True
    pc0.font.color.rgb = CYAN_ACCENT
    pc0.space_after = Pt(14)

    c_points = [
        ("Multi-Network Profile Isolation: Automatically detects SSID changes, isolates device tables per Wi-Fi environment, and offers seamless cross-network switching.",),
        ("Zero-Configuration Discovery: Automatically sweeps local subnets and provides targeted unicast (⚡ Probe IP) for custom network nodes.",),
        ("Hardware Brand Recognition: Resolves manufacturer details (Apple, Samsung, Intel, etc.) directly from hardware MAC OUI prefixes.",),
        ("Auditable DB Persistence: Permanently stores all connection/disconnection logs in SQLite with 1-click CSV report exports.",),
        ("Enterprise Governance & Live Alerts: Secure dual-role RBAC authentication combined with sub-second WebSocket toast notifications.",)
    ]
    c_points = [p[0] for p in c_points]

    for pt in c_points:
        p = tf_c.add_paragraph()
        p.text = "✔  " + pt
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(10)

    # Save presentation
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == '__main__':
    output_dir = r"C:\Users\hp\.gemini\antigravity-ide\scratch\network-device-scanner"
    output_file = os.path.join(output_dir, "NetScan_Pro_Presentation.pptx")
    create_presentation(output_file)
